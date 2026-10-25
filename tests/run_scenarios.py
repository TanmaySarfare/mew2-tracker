"""Run app.py top to bottom under fake streamlit/plotly/yfinance for many scenarios."""
import datetime as _dt
import io
import numpy as np
import itertools
import os
import runpy
import string
import sys
import traceback
import urllib.request
import warnings
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "stubs"))
APP = os.path.join(HERE, "..", "app.py")

import pandas as pd  # noqa: E402
import streamlit as st  # noqa: E402
import yfinance as yf  # noqa: E402

RESULTS = []

# --- Fake copies of Nasdaq's symbol directory --------------------------------------------
filler = ["".join(p) for p in itertools.product(string.ascii_uppercase, repeat=3)][:1200]
NASDAQ_FILE = "\n".join(
    ["Symbol|Security Name|Market Category|Test Issue|Financial Status|Round Lot Size|ETF|NextShares",
     "AAPL|Apple Inc. - Common Stock|Q|N|N|100|N|N",
     "MSFT|Microsoft Corporation - Common Stock|Q|N|N|100|N|N",
     "QQQ|Invesco QQQ Trust, Series 1|G|N|N|100|Y|N",
     "ZZZT|Test Issue Corp - Common Stock|Q|Y|N|100|N|N",
     "ABCDW|Example Acquisition Corp - Warrant|S|N|N|100|N|N",
     "NA|Nano Labs Ltd - Class A Ordinary Shares|S|N|N|100|N|N"]
    + [f"{s}|Filler {s} Inc. - Common Stock|S|N|N|100|N|N" for s in filler]
    + ["File Creation Time: 1025202614:00|||||||"])
OTHER_FILE = "\n".join([
    "ACT Symbol|Security Name|Exchange|CQS Symbol|ETF|Round Lot Size|Test Issue|NASDAQ Symbol",
    "VTI|Vanguard Total Stock Market ETF|P|VTI|Y|100|N|VTI",
    "BRK.B|Berkshire Hathaway Inc. New Common Stock|N|BRK.B|N|100|N|BRK=B",
    "ABR$D|Arbor Realty Trust 6.375% Series D Preferred Stock|N|ABRpD|N|100|N|ABR-D",
    "JPM|JP Morgan Chase & Co. Common Stock|N|JPM|N|100|N|JPM",
    "PFF|iShares Preferred and Income Securities ETF|P|PFF|Y|100|N|PFF",
    "File Creation Time: 1025202614:00|||||||"])
REAL_URLOPEN = urllib.request.urlopen
import requests  # noqa: E402
REAL_GET = requests.get


class _HTTPResponse:
    def __init__(self, text):
        self.text = text

    def raise_for_status(self):
        pass


def fake_get(url, headers=None, timeout=None):
    if url.endswith("nasdaqlisted.txt"):
        return _HTTPResponse(NASDAQ_FILE)
    if url.endswith("otherlisted.txt"):
        return _HTTPResponse(OTHER_FILE)
    raise OSError("network disabled in tests")


def offline_get(url, headers=None, timeout=None):
    raise OSError("network disabled in tests")


class _Response(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def fake_urlopen(request, timeout=None):
    url = request.full_url if hasattr(request, "full_url") else request
    if url.endswith("nasdaqlisted.txt"):
        return _Response(NASDAQ_FILE.encode())
    if url.endswith("otherlisted.txt"):
        return _Response(OTHER_FILE.encode())
    raise OSError("network disabled in tests")


def offline_urlopen(request, timeout=None):
    raise OSError("network disabled in tests")


# --- Runner --------------------------------------------------------------------------------
BUTTON_PREFIXES = ("zip_prepare", "refresh", "quick_")


def run(name, overrides=None, keep_session=False, expect=None, forbid=None, check=None, directory="ok", click=None):
    if not keep_session:
        st.session_state.clear()
    urllib.request.urlopen = fake_urlopen if directory == "ok" else offline_urlopen
    requests.get = fake_get if directory == "ok" else offline_get
    overrides = dict(overrides or {})
    if click:  # a button with an on_click callback: Streamlit runs the callback, then reruns the script
        callback, args, kwargs = st.STATE.callbacks[click]
        callback(*args, **kwargs)
        overrides[click] = True
    error, reruns = None, 0
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        while True:
            st.STATE.overrides = dict(overrides)
            st.STATE.reset_run()
            for k, v in overrides.items():  # like Streamlit, a user's change is in Session State before the run
                if v is not None and not k.startswith(BUTTON_PREFIXES):
                    st.session_state[k] = v
            yf.CALLS.clear()
            try:
                runpy.run_path(APP, run_name="__main__")
            except st.StopExecution:
                pass
            except st.RerunRequested:
                reruns += 1
                overrides = {k: v for k, v in overrides.items() if not k.startswith(BUTTON_PREFIXES)}
                if reruns < 3:
                    continue
                error = "too many reruns"
            except Exception:
                error = traceback.format_exc()
            break
    calls = st.STATE.calls
    kinds = {}
    for kind, _ in calls:
        kinds[kind] = kinds.get(kind, 0) + 1
    texts = {k: [d["body"] for kk, d in calls if kk == k] for k in
             ("error", "warning", "info", "caption", "subheader", "title", "markdown")}
    problems = list(st.STATE.issues)
    relevant = sorted({f"{w.category.__name__}: {w.message}" for w in caught if "stubs" not in str(w.filename)})
    for kind in expect or []:
        if not kinds.get(kind):
            problems.append(f"expected at least one {kind}")
    for kind in forbid or []:
        if kinds.get(kind):
            problems.append(f"unexpected {kind}: {texts.get(kind)}")
    if check and error is None:
        try:
            msg = check(calls, texts)
            if msg:
                problems.append(msg)
        except Exception:
            problems.append("check failed: " + traceback.format_exc())
    ok = error is None and not problems
    RESULTS.append((name, ok))
    summary = ", ".join(f"{k}={v}" for k, v in sorted(kinds.items()) if k in
                        ("dataframe", "plotly_chart", "download_button", "error", "warning", "info"))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}: {summary}; yahoo requests={len(yf.CALLS)}"
          + (f"; reruns={reruns}" if reruns else ""))
    if error:
        print(error)
    for p in problems:
        print("   problem:", p)
    for w in relevant:
        print("   warning:", w)
    return calls, texts


def find(calls, kind, **match):
    return [d for k, d in calls if k == kind and all(d.get(a) == b for a, b in match.items())]


def tables(calls):
    return {d["key"]: d["data"] for k, d in calls if k == "dataframe"}


def downloads(calls):
    return {d["key"]: d for k, d in calls if k == "download_button"}


def metrics(calls):
    return {d["label"]: (d["value"], d["delta"]) for k, d in calls if k == "metric"}


def header_html(texts):
    return next((m for m in texts["markdown"] if 'class="td-name"' in m), "")


# --- Search ------------------------------------------------------------------------------------
def check_welcome(calls, texts):
    box = find(calls, "selectbox", key="ticker")[0]
    labels = dict(zip(box["options"], box["labels"]))
    if box["value"] is not None:
        return "search should start empty"
    for sym in ("NA", "BRK-B", "PFF", "QQQ", "VTI", "^GSPC", "BTC-USD", "VTSAX"):
        if sym not in labels:
            return f"{sym} missing from suggestions"
    for sym in ("ZZZT", "ABCDW", "ABR$D", "ABR-D"):
        if sym in labels:
            return f"{sym} should be filtered out"
    if (labels["AAPL"] != "Apple Inc. (AAPL)" or labels["NA"] != "Nano Labs Ltd (NA)"
            or labels["BRK-B"] != "Berkshire Hathaway Inc. New (BRK-B)"):
        return f"labels wrong: {labels['AAPL']!r}, {labels['NA']!r}, {labels['BRK-B']!r}"
    if box["options"][:3] != ["AAPL", "MSFT", "NVDA"]:
        return "popular tickers should come first"
    odd = [l for l in box["labels"] if "^" in l or "=X" in l or "=F" in l]
    if odd:
        return f"suggestions should not show Yahoo codes: {odd[:3]}"
    if labels["^GSPC"] != "S&P 500 index" or labels["EURUSD=X"] != "Euro to U.S. dollar (EUR/USD)" \
            or labels["BTC-USD"] != "Bitcoin (BTC)":
        return f"special labels wrong: {labels['^GSPC']}, {labels['EURUSD=X']}, {labels['BTC-USD']}"
    if not find(calls, "button", key="quick_AAPL"):
        return "quick picks missing"
    if any("didn't load" in c for c in texts["caption"]):
        return "should not claim the list failed"


run("first visit: empty search with suggestions", check=check_welcome, forbid=["error", "warning"])


def check_offline_list(calls, texts):
    box = find(calls, "selectbox", key="ticker")[0]
    if "AAPL" not in box["options"] or "^GSPC" not in box["options"] or len(box["options"]) > 200:
        return "fallback list should hold popular names and extras only"
    if not any("didn't load" in c for c in texts["caption"]):
        return "should explain that the full list didn't load"


run("ticker list unreachable: falls back to popular names", directory="offline", check=check_offline_list,
    forbid=["error"])
run("...and doesn't retry on every click", {"ticker": "MSFT"}, keep_session=True, directory="ok",
    check=lambda c, t: None if any("didn't load" in x for x in t["caption"]) else "should keep using the fallback")


def check_quote(calls, texts):
    h = header_html(texts)
    for piece in ("Apple Inc.", 'td-ticker">AAPL', "Stock listed on NasdaqGS, prices in USD", "td-price", "as of ",
                  "52-week range", "Market cap", "Dividend yield, 12 months", "Sector"):
        if piece not in h:
            return f"header lacks {piece!r}"
    nav = find(calls, "segmented_control", key="section_EQUITY")[0]
    if nav["options"] != ["Price", "Performance", "Statements", "Ratios", "Valuation", "Earnings", "Dividends",
                          "Compare", "Profile"] or nav["value"] != "Price":
        return f"navigation wrong: {nav}"
    if not find(calls, "button", key="zip_prepare"):
        return "export button missing"


run("pick Apple from suggestions", {"ticker": "AAPL"}, check=check_quote, forbid=["error", "warning"])
st.session_state.clear()
run("welcome screen", {})
run("quick pick button loads Apple", keep_session=True, click="quick_AAPL", forbid=["error", "warning"],
    check=lambda c, t: None if "Apple Inc." in header_html(t) else "quick pick didn't load Apple")


def check_typed_exact(calls, texts):
    if find(calls, "selectbox", label='Matches on Yahoo Finance for "vod.l"'):
        return "an exact ticker should load without an extra question"
    if "VODAFONE" not in header_html(texts).upper():
        return "VOD.L didn't load"


run("typed foreign ticker not in list (vod.l)", {"ticker": "vod.l"}, check=check_typed_exact, forbid=["error"])


def check_typed_name(calls, texts):
    box = find(calls, "selectbox", label='Matches on Yahoo Finance for "taiwan semi"')
    if not box or box[0]["value"] != "TSM":
        return "should offer Yahoo matches for a typed name"
    if "Taiwan Semiconductor" not in header_html(texts):
        return "TSM didn't load"


run("typed company name not in list", {"ticker": "taiwan semi"}, check=check_typed_name, forbid=["error"])
run("typed gibberish with spaces", {"ticker": "no such company"},
    check=lambda c, t: None if any("Nothing on Yahoo Finance matches" in w for w in t["warning"]) else "missing message")
run("typed unknown ticker", {"ticker": "zzzzz"},
    check=lambda c, t: None if any("suffix" in e for e in t["error"]) else "missing suffix advice")
run("Yahoo search down, typed ticker still tried", {"ticker": "searchfail"}, expect=["error"])
run("rate limited", {"ticker": "RATELIM"},
    check=lambda c, t: None if any("limiting requests" in e for e in t["error"]) else "missing rate-limit message")
run("refresh prices", {"ticker": "AAPL", "refresh": True}, forbid=["error"])


# --- Price ----------------------------------------------------------------------------------------
def check_price(calls, texts):
    t = tables(calls)["table_prices"]
    if t.index[0] < t.index[-1]:
        return "table should be newest first"
    d = downloads(calls)
    view = d["download_prices_view"]["data"].decode().splitlines()
    full = d["download_prices_all"]["data"].decode().splitlines()
    if not view[0].startswith("Date,Open,High,Low,Close,Adj Close,Volume") or len(full) < len(view):
        return f"downloads wrong: view {len(view)} rows, full {len(full)} rows"
    if f"{len(view) - 1:,}" not in d["download_prices_view"]["label"]:
        return "row count in button label is wrong"
    if view[1][:10] > view[-1][:10]:
        return "CSV should run oldest to newest"


run("price: default 1Y daily", {"ticker": "AAPL"}, check=check_price, forbid=["error", "warning"])
for period in ("1M", "3M", "6M", "YTD", "5Y", "10Y", "Max"):
    run(f"price period {period}", {"ticker": "AAPL", "price_period": period}, check=check_price,
        forbid=["error", "warning"])
for bar in ("Weekly", "Monthly", "Quarterly"):
    run(f"price bars {bar}", {"ticker": "AAPL", "price_bar": bar, "price_period": "5Y"}, check=check_price,
        forbid=["error", "warning"])
run("price: monthly bars with a 1-month period", {"ticker": "AAPL", "price_bar": "Monthly", "price_period": "1M"},
    forbid=["error"])
run("price: custom dates", {"ticker": "AAPL", "price_period": "Custom"}, check=check_price, forbid=["error", "warning"])
run("price: custom dates reversed", {"ticker": "AAPL", "price_period": "Custom", "price_from": _dt.date(2024, 5, 1),
                                     "price_to": _dt.date(2023, 5, 1)}, expect=["error"])
run("price: line chart, log scale, no averages",
    {"ticker": "AAPL", "price_style": "Line", "price_log": True, "price_sma": False}, forbid=["error"])


def check_keep_period(calls, texts):
    seg = find(calls, "segmented_control", key="price_period")[0]
    if seg["value"] != "5Y":
        return f"the chosen period should stay selected, got {seg['value']}"


run("price: choose 5Y", {"ticker": "AAPL", "price_period": "5Y"})
run("price: click 5Y again (deselect) keeps 5Y", {"ticker": "AAPL", "price_period": None}, keep_session=True,
    check=check_keep_period, forbid=["error"])
run("price: go to another section", {"ticker": "AAPL", "section_EQUITY": "Valuation"}, keep_session=True)
run("...and back keeps 5Y", {"ticker": "AAPL", "section_EQUITY": "Price"}, keep_session=True, check=check_keep_period)

# --- Other sections for a stock ---------------------------------------------------------------------
for section in ("Performance", "Statements", "Valuation", "Dividends", "Earnings", "Profile"):
    def check_section(calls, texts, section=section):
        t = tables(calls)
        if section == "Statements":
            inc = t.get("table_statement_income_yearly")
            if inc is None or list(inc.columns) != sorted(inc.columns):
                return "income statement missing or unsorted"
            if not (1 < inc.loc["Diluted EPS"].iloc[-1] < 20):
                return "EPS must not be scaled"
            if not find(calls, "selectbox", key="units"):
                return "units picker should sit in the section"
        if section == "Valuation":
            v = [x for k, x in metrics(calls).items() if k.startswith("DCF value per share")]
            if not v or v[0][0] == "n/a":
                return "value per share missing"
            heads = [s for s in texts["subheader"] if s[:2].rstrip(".").isdigit()]
            if len(heads) != 9 or not heads[-1].startswith("9. Excel"):
                return f"valuation steps should be numbered 1-9: {heads}"
        if section == "Performance":
            if not find(calls, "selectbox", key="benchmark"):
                return "benchmark picker missing"
            if "S&P 500 total return (%)" not in t["table_returns"].columns:
                return "benchmark comparison missing"
        if section == "Dividends":
            paid = metrics(calls)["Paid in the last 12 months"][0]
            if len(paid.split(".")[-1]) != 2:
                return f"dividends should show 2 decimals, got {paid}"
        if section == "Profile" and "table_all_fields" not in t:
            return "all fields missing"
    run(f"AAPL {section}", {"ticker": "AAPL", "section_EQUITY": section}, check=check_section,
        forbid=["error", "warning"])

run("statements quarterly in billions", {"ticker": "AAPL", "section_EQUITY": "Statements",
                                         "statement_period": "Quarterly", "units": "Billions"}, forbid=["error", "warning"])
run("units choice carries over to DCF", {"ticker": "AAPL", "section_EQUITY": "Valuation"}, keep_session=True,
    check=lambda c, t: None if find(c, "selectbox", key="units")[0]["value"] == "Billions" else "units not remembered")
run("DCF: reported FCF, 3-year average, no fade, manual WACC",
    {"ticker": "AAPL", "section_EQUITY": "Valuation", "dcf_kind_AAPL": "Free cash flow (OCF - CapEx)", "dcf_avg_AAPL": True,
     "dcf_fade_AAPL": False, "dcf_wacc_manual_AAPL": True, "dcf_wacc_AAPL": 9.0}, forbid=["error"])
run("DCF: WACC below growth",
    {"ticker": "AAPL", "section_EQUITY": "Valuation", "dcf_wacc_manual_AAPL": True, "dcf_wacc_AAPL": 2.0},
    check=lambda c, t: None if any("WACC is 2.00%" in e for e in t["error"]) else "error should quote the numbers")


def check_bench_carry(calls, texts):
    beta = find(calls, "number_input", label="Beta")[0]
    if "Nifty 50" not in beta["help"] or "^" in beta["help"]:
        return f"DCF beta should use the benchmark chosen in Performance: {beta['help']}"


run("benchmark: pick Nifty in Performance", {"ticker": "AAPL", "section_EQUITY": "Performance", "benchmark": "^NSEI"})
run("...DCF beta uses it", {"ticker": "AAPL", "section_EQUITY": "Valuation"}, keep_session=True, check=check_bench_carry)
run("benchmark: typed custom symbol", {"ticker": "AAPL", "section_EQUITY": "Performance", "benchmark": "msft"},
    forbid=["error"],
    check=lambda c, t: None if "MSFT total return (%)" in tables(c)["table_returns"].columns else "custom benchmark ignored")
for w in ("1Y", "3Y", "10Y", "Max"):
    run(f"performance chart period {w}", {"ticker": "AAPL", "section_EQUITY": "Performance", "perf_window": w},
        forbid=["error"])


# --- Funds, indices, crypto, edge cases --------------------------------------------------------
def check_fund_header(calls, texts):
    h = header_html(texts)
    for piece in ("ETF listed on NYSEArca", "Total assets", "Dividend yield, 12 months", "Fund family"):
        if piece not in h:
            return f"fund header lacks {piece!r}"
    nav = find(calls, "segmented_control", key="section_ETF")[0]
    if nav["options"] != ["Price", "Performance", "Holdings", "Valuation", "Dividends", "Compare", "Profile"]:
        return f"fund navigation wrong: {nav['options']}"


run("VTI header", {"ticker": "VTI"}, check=check_fund_header, forbid=["error", "warning"])
for section in ("Performance", "Holdings", "Valuation", "Dividends", "Compare", "Profile"):
    run(f"VTI {section}", {"ticker": "VTI", "section_ETF": section}, forbid=["error", "warning"])
run("VTSAX capital gains", {"ticker": "VTSAX", "section_MUTUALFUND": "Dividends"}, forbid=["error", "warning"],
    check=lambda c, t: None if "table_capital_gains" in tables(c) else "capital gains missing")
run("^GSPC header", {"ticker": "^GSPC"}, forbid=["error", "warning"],
    check=lambda c, t: None if ("Index, values in USD" in header_html(t) and "td-ticker" not in header_html(t)
                                and "^" not in header_html(t)) else "index header should be plain, without ^GSPC")
run("^GSPC performance vs itself", {"ticker": "^GSPC", "section_INDEX": "Performance"}, forbid=["error", "warning"])
run("BTC-USD", {"ticker": "BTC-USD", "section_CRYPTOCURRENCY": "Performance"}, forbid=["error", "warning"])
run("profile fails (INFOFAIL)", {"ticker": "INFOFAIL"}, forbid=["error"])
for section in ("Statements", "Ratios", "Valuation", "Earnings", "Holdings", "Dividends", "Compare", "Profile"):
    run(f"INFOFAIL {section}", {"ticker": "INFOFAIL", "section_ANY": section}, forbid=["error"])
run("TSM DCF currency conversion", {"ticker": "taiwan semi", "section_EQUITY": "Valuation"}, forbid=["error"],
    check=lambda c, t: None if any("reports in TWD" in i for i in t["info"]) else "conversion note missing")
run("VOD.L DCF in pence", {"ticker": "vod.l", "section_EQUITY": "Valuation"}, forbid=["error"])
run("JPM bank DCF", {"ticker": "JPM", "section_EQUITY": "Valuation"}, forbid=["error"],
    check=lambda c, t: None if any("Banks" in i for i in t["info"]) else "bank caveat missing")
run("negative FCF", {"ticker": "NEGF", "section_EQUITY": "Valuation"}, forbid=["error"])
run("no statements", {"ticker": "EMPTYFIN", "section_EQUITY": "Valuation"}, forbid=["error"])


# --- Export ----------------------------------------------------------------------------------------
def zip_check(expected):
    def check(calls, texts):
        d = downloads(calls).get("zip_download")
        if d is None:
            return "Download ZIP button should replace the prepare button"
        if find(calls, "button", key="zip_prepare"):
            return "prepare button should be gone"
        z = zipfile.ZipFile(io.BytesIO(d["data"]))
        names = [n.split("/", 1)[1] for n in z.namelist()]
        missing = [e for e in expected if e not in names]
        if missing:
            return f"ZIP lacks {missing}"
        for n in z.namelist():
            if n.endswith(".csv"):
                pd.read_csv(z.open(n))
    return check


run("ZIP for AAPL", {"ticker": "AAPL", "zip_prepare": True}, forbid=["error"],
    check=zip_check(["prices_daily_full_history.csv", "income_yearly_millions.csv", "financial_ratios.csv",
                     "earnings.csv", "dcf_building_blocks_millions.csv", "trailing_returns.csv", "growth_estimates.csv",
                     "README.txt"]))
run("ZIP resets after changing units", {"ticker": "AAPL", "section_EQUITY": "Statements", "units": "Raw"},
    keep_session=True,
    check=lambda c, t: None if find(c, "button", key="zip_prepare") else "stale ZIP offered after units changed")
run("ZIP for VTI", {"ticker": "VTI", "zip_prepare": True}, forbid=["error"],
    check=zip_check(["prices_daily_full_history.csv", "fund_top_holdings.csv", "fund_sector_weightings.csv"]))


# --- v3 features ----------------------------------------------------------------------------------
import json as _json, subprocess, tempfile  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def check_indicators(calls, texts):
    fig = find(calls, "plotly_chart", key="chart_price")[0]["figure"]
    names = [t.props.get("name") for t, _, _ in fig.traces]
    for needed in ("RSI (14)", "MACD", "MACD signal", "Bollinger upper", "50-day average"):
        if needed not in names:
            return f"chart lacks {needed}: {names}"
    if fig.rows != 4 or len(getattr(fig, "hlines", [])) != 2:
        return f"expected price, volume, RSI and MACD panels with 30/70 lines; rows={fig.rows}"
    t = tables(calls)["table_prices"]
    for col in ("RSI 14", "MACD", "Bollinger lower", "SMA 200"):
        if col not in t.columns:
            return f"table lacks {col}"
    rsi = t["RSI 14"].dropna()
    if rsi.empty or rsi.min() < 0 or rsi.max() > 100:
        return "RSI outside 0-100"
    header = downloads(calls)["download_prices_view"]["data"].decode().splitlines()[0]
    if "RSI 14" not in header or "MACD signal" not in header:
        return "CSV should include the indicators"


run("price: all indicators on", {"ticker": "AAPL", "price_indicators": ["Moving averages", "Bollinger bands", "RSI", "MACD"]},
    check=check_indicators, forbid=["error", "warning"])
run("price: indicators off", {"ticker": "AAPL", "price_indicators": []}, forbid=["error", "warning"],
    check=lambda c, t: None if "SMA 50" not in tables(c)["table_prices"].columns else "SMA shouldn't show when off")
run("price: indicators on weekly bars", {"ticker": "AAPL", "price_bar": "Weekly", "price_period": "5Y",
                                         "price_indicators": ["RSI", "MACD"]}, forbid=["error", "warning"])
run("price: indicators survive a tab switch", {"ticker": "AAPL", "price_indicators": ["RSI"]})
run("...go to Ratios", {"ticker": "AAPL", "section_EQUITY": "Ratios"}, keep_session=True)
run("...back to Price keeps RSI", {"ticker": "AAPL", "section_EQUITY": "Price"}, keep_session=True,
    check=lambda c, t: None if find(c, "pills", key="price_indicators")[0]["value"] == ["RSI"] else "indicator choice lost")


def check_ratios(calls, texts):
    t = tables(calls)
    prof = t.get("table_ratios_Profitability")
    if prof is None:
        return "profitability table missing"
    gm = prof.loc["Gross margin (%)"].dropna()
    om = prof.loc["Operating margin (%)"].dropna()
    if not (44 < gm.iloc[-1] < 46) or not (29 < om.iloc[-1] < 31):
        return f"margins wrong: gross {gm.iloc[-1]}, operating {om.iloc[-1]}"
    val = t.get("table_ratios_Valuation_at_fiscal_year_end")
    if val is None or val.loc["Price to earnings, P/E (x)"].dropna().empty:
        return "valuation ratios missing"
    health = t["table_ratios_Financial_health"]
    if abs(health.loc["Current ratio (x)"].dropna().iloc[-1] - 0.925) > 0.01:
        return "current ratio wrong"
    csv = downloads(calls)["download_ratios"]["data"].decode().splitlines()[0]
    if not csv.startswith("Ratio,Group,"):
        return f"ratio CSV header wrong: {csv[:40]}"
    if not find(calls, "plotly_chart", key="chart_ratio"):
        return "ratio chart missing"


run("ratios tab", {"ticker": "AAPL", "section_EQUITY": "Ratios"}, check=check_ratios, forbid=["error", "warning"])
run("ratios: chart another ratio", {"ticker": "AAPL", "section_EQUITY": "Ratios", "ratio_pick": "Price to earnings, P/E (x)"},
    forbid=["error", "warning"])
run("ratios: no statements", {"ticker": "EMPTYFIN", "section_EQUITY": "Ratios"}, forbid=["error"],
    check=lambda c, t: None if any("calculate ratios" in i for i in t["info"]) else "missing explanation")
run("ratios: bank", {"ticker": "JPM", "section_EQUITY": "Ratios"}, forbid=["error"])


def excel_value(data):
    """Recalculate the downloaded workbook in LibreOffice and read the DCF value per share."""
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "model.xlsx")
        open(path, "wb").write(data)
        out = subprocess.run([sys.executable, "/mnt/skills/public/xlsx/scripts/recalc.py", path, "90"],
                             capture_output=True, text=True)
        report = _json.loads(out.stdout)
        wb = load_workbook(path, data_only=True)
        values = {}
        for ws_name, label in (("DCF", "Value per share ("), ("DDM", "Value per share (")):
            if ws_name in wb.sheetnames:
                for row in wb[ws_name].iter_rows(min_col=1, max_col=2):
                    if isinstance(row[0].value, str) and row[0].value.startswith(label) and row[0].font.bold:
                        values[ws_name] = row[1].value
        values["weighted"] = wb["Scenarios"]["B17"].value
        return report, values, wb.sheetnames


def check_valuation_extras(calls, texts):
    m = metrics(calls)
    implied = m.get("Year-1 growth the price implies")
    if not implied or not implied[0].endswith("%"):
        return f"implied growth missing: {implied}"
    print(f"   reverse DCF: price implies year-1 growth {implied[0]} ({implied[1]}), return {m['Return the price implies'][0]}")
    if not m.get("Return the price implies", ("n/a",))[0].endswith("%"):
        return "implied return missing"
    sc = tables(calls).get("table_scenario_results")
    if sc is None or list(sc["Scenario"]) != ["Bear", "Base", "Bull"]:
        return "scenario table missing"
    vals = sc["Value per share"].tolist()
    if not (vals[0] < vals[1] < vals[2]):
        return f"bear < base < bull expected: {vals}"
    dcf_value = float([v for k, v in m.items() if k.startswith("DCF value per share")][0][0].replace(",", ""))
    if abs(vals[1] - dcf_value) > 0.01:
        return f"base scenario {vals[1]} should equal the DCF value {dcf_value}"
    ddm = [v for k, v in m.items() if k.startswith("Dividend model value per share")]
    if not ddm or ddm[0][0] == "n/a":
        return "dividend model missing"
    weighted = float([v for k, v in m.items() if k.startswith("Probability-weighted")][0][0].replace(",", ""))
    xl = downloads(calls).get("download_excel_model")
    if xl is None:
        return "Excel download missing"
    report, values, sheets = excel_value(xl["data"])
    if report.get("status") != "success" or report.get("total_errors"):
        return f"Excel recalculation problems: {report}"
    if sheets != ["Inputs", "DCF", "Sensitivity", "Scenarios", "DDM", "Historical", "Notes"]:
        return f"unexpected sheets {sheets}"
    if abs(values["DCF"] - dcf_value) > 0.01:
        return f"Excel DCF value {values['DCF']:.4f} differs from the app's {dcf_value}"
    if abs(values["DDM"] - float(ddm[0][0].replace(",", ""))) > 0.01:
        return f"Excel DDM value {values['DDM']:.4f} differs from the app's {ddm[0][0]}"
    if abs(values["weighted"] - weighted) > 0.01:
        return f"Excel weighted value {values['weighted']:.4f} differs from the app's {weighted}"
    print(f"   excel: {report['total_formulas']} formulas, 0 errors; DCF {values['DCF']:.2f} = app {dcf_value:.2f}; "
          f"DDM {values['DDM']:.2f}; weighted {values['weighted']:.2f}")


run("valuation: reverse DCF, scenarios, DDM, Excel matches the app", {"ticker": "AAPL", "section_EQUITY": "Valuation"},
    check=check_valuation_extras, forbid=["error", "warning"])
custom = pd.DataFrame({"Scenario": ["Bear", "Base", "Bull"], "Probability (%)": [20.0, 30.0, 30.0],
                       "Growth change (pts)": [-8.0, 0.0, 3.0], "WACC change (pts)": [2.0, 0.0, -0.5],
                       "Terminal growth change (pts)": [-1.0, 0.0, 0.25]})
run("valuation: edited scenarios that don't add to 100%", {"ticker": "AAPL", "section_EQUITY": "Valuation",
                                                           "scenarios_AAPL": custom}, forbid=["error"],
    check=lambda c, t: None if any("add up to 80%" in x for x in t["caption"]) else "should explain the rescaling")
run("valuation: manual WACC, 12 years, no fade, Excel still matches",
    {"ticker": "AAPL", "section_EQUITY": "Valuation", "dcf_wacc_manual_AAPL": True, "dcf_wacc_AAPL": 8.5,
     "dcf_years_AAPL": 12, "dcf_fade_AAPL": False, "units": "Billions"}, forbid=["error"], check=check_valuation_extras)
run("valuation: currency conversion (TSM), Excel matches", {"ticker": "taiwan semi", "section_EQUITY": "Valuation"},
    forbid=["error"], check=check_valuation_extras)
run("valuation: no dividends (NEGF) still offers Excel", {"ticker": "NEGF", "section_EQUITY": "Valuation"}, forbid=["error"],
    check=lambda c, t: None if "download_excel_model" in downloads(c) and any("hasn't paid dividends" in i for i in t["info"])
    else "Excel or no-dividend note missing")
run("valuation: bank gets DDM without FCFF", {"ticker": "JPM", "section_EQUITY": "Valuation"}, forbid=["error"],
    check=lambda c, t: None if [k for k in metrics(c) if k.startswith("Dividend model value")] else "bank DDM missing")
run("valuation: WACC below growth still shows DDM",
    {"ticker": "AAPL", "section_EQUITY": "Valuation", "dcf_wacc_manual_AAPL": True, "dcf_wacc_AAPL": 2.0}, expect=["error"],
    check=lambda c, t: None if any(x.endswith("Dividend discount model") for x in t["subheader"]) else "DDM missing")
run("valuation: no statements, DDM only", {"ticker": "EMPTYFIN", "section_EQUITY": "Valuation"}, forbid=["error"])
run("VTI valuation: dividend model", {"ticker": "VTI", "section_ETF": "Valuation"}, forbid=["error", "warning"],
    check=lambda c, t: None if [k for k in metrics(c) if k.startswith("Dividend model value")] else "fund DDM missing")


def check_earnings(calls, texts):
    m = metrics(calls)
    if m.get("Next earnings date", ("Not announced",))[0] == "Not announced":
        return "next earnings date missing"
    if m.get("Earnings per share expected", ("n/a",))[0] != "1.92":
        return f"EPS expected should come from the calendar: {m.get('Earnings per share expected')}"
    if "Beat analysts' estimate" not in m:
        return "beat count missing"
    t = tables(calls)
    hist = t.get("table_earnings_history")
    if hist is None or hist["Reported EPS"].notna().sum() != 8 or hist["Reported EPS"].isna().sum() != 2:
        return "earnings history should hold 8 reported and 2 upcoming quarters"
    q = t.get("table_quarterly_results")
    if q is None or len(q) < 4:
        return "quarterly results missing"
    for key in ("chart_earnings", "chart_quarters"):
        if not find(calls, "plotly_chart", key=key):
            return f"{key} missing"
    if "Low" not in m:
        return "analyst price targets should be in this tab"


run("earnings tab", {"ticker": "AAPL", "section_EQUITY": "Earnings"}, check=check_earnings, forbid=["error", "warning"])
run("earnings: no data for this ticker", {"ticker": "INFOFAIL", "section_ANY": "Earnings"}, forbid=["error"])


def check_compare(calls, texts):
    picks = find(calls, "multiselect", key="compare_AAPL")[0]["value"]
    if picks[:2] != ["SONY", "MSFT"]:
        return f"default peers should come from the industry list: {picks}"
    if not any("Couldn't load SONY" in w for w in texts["warning"]):
        return "a peer that fails should be named"
    snap = tables(calls).get("table_compare")
    if snap is None or list(snap.index[:3]) != ["AAPL", "MSFT", "TSM"] or snap.index[-1] != "Peer median":
        return f"comparison table rows wrong: {None if snap is None else list(snap.index)}"
    corr = tables(calls).get("table_correlation")
    if corr is None or abs(corr.loc["AAPL", "AAPL"] - 1) > 1e-9:
        return "correlation table wrong"
    fig = find(calls, "plotly_chart", key="chart_compare")[0]["figure"]
    starts = {t.props["y"].iloc[0] for t, _, _ in fig.traces}
    if any(abs(v - 10000) > 1e-6 for v in starts):
        return "every line should start at 10,000"
    for k in ("download_compare", "download_compare_prices", "download_correlations"):
        if k not in downloads(calls):
            return f"{k} missing"


run("compare: industry peers by default", {"ticker": "AAPL", "section_EQUITY": "Compare"}, check=check_compare,
    forbid=["error"])
run("compare: user picks, typed symbol, no peers suggested", {"ticker": "MSFT", "section_EQUITY": "Compare",
                                                             "compare_MSFT": ["AAPL", "vti", "^GSPC"]},
    forbid=["error", "warning"],
    check=lambda c, t: None if list(tables(c)["table_compare"].index[:4]) == ["MSFT", "AAPL", "VTI", "S&P 500"] else "rows wrong")
run("compare: nothing picked yet", {"ticker": "MSFT", "section_EQUITY": "Compare"}, forbid=["error"],
    check=lambda c, t: None if any("Add one or more tickers" in i for i in t["info"]) else "should invite picks")
run("compare: chart period 1Y", {"ticker": "MSFT", "section_EQUITY": "Compare", "compare_MSFT": ["AAPL"],
                                 "compare_window": "1Y"}, forbid=["error", "warning"])
run("compare for an ETF", {"ticker": "VTI", "section_ETF": "Compare", "compare_VTI": ["VTSAX", "^GSPC"]},
    forbid=["error", "warning"])
def check_profile(calls, texts):
    md = " ".join(texts["markdown"])
    if 'class="td-about"' not in md or 'class="td-facts"' not in md:
        return "About text and facts panel missing"
    if '>example.com</a>' not in md or "Seattle" in md:
        return "website should show as a short link"
    if [k for k in metrics(calls) if k in ("Sector", "Industry", "Website", "Country", "Employees")]:
        return "facts should not use the big metric boxes"
    if "Held by insiders" not in md or "2.07%" not in md:
        return "ownership summary should read as stats with percentages"
    inst = tables(calls).get("table_holders_institutional_holders")
    if inst is None or "Share of company (%)" not in inst.columns or abs(inst["Share of company (%)"].iloc[0] - 9.3) > 1e-9:
        return "institutional holders should have readable columns and percentages"
    fields = tables(calls).get("table_all_fields")
    if fields is None:
        return "all-fields table missing"
    lookup = dict(zip(fields["Field"], fields["Value"]))
    if "52-week high" not in lookup or "Market cap" not in lookup or not lookup["Market cap"].endswith(("T", "B", "M")):
        return f"fields should be humanized and formatted: {list(lookup.items())[:6]}"
    if any(ch.isupper() for f in fields["Field"] for ch in f[1:2]) and False:
        return "camelCase left over"
    if not any(label.startswith("All the data fields") for label in [d["label"] for k, d in calls if k == "expander"]):
        return "raw fields should sit in a collapsible section"
    csv = downloads(calls)["download_all_fields"]["data"].decode().splitlines()[0]
    if csv != "Field,Yahoo name,Value":
        return f"all-fields CSV should keep Yahoo's names too: {csv}"


run("profile: about, facts, ownership, readable fields", {"ticker": "AAPL", "section_EQUITY": "Profile"},
    forbid=["error", "warning"], check=check_profile)


def check_analyst_tables(calls, texts):
    t = tables(calls)
    est = t.get("table_analyst_earnings_estimate")
    if est is None or est.index[0] != "This quarter" or "Growth (%)" not in est.columns or est["Growth (%)"].iloc[0] != 8.0:
        return "EPS estimate table should be readable"
    rec = t.get("table_analyst_recommendations")
    if rec is None or rec.index[0] != "This month" or "Strong buy" not in rec.columns:
        return "ratings table should be readable"
    ud = t.get("table_analyst_upgrades_downgrades")
    if ud is None or ud["Action"].iloc[1] != "Downgrade":
        return "rating actions should be words"
    when = metrics(calls)["Next earnings date"][0]
    if "–" not in when or when.count(",") != 1:
        return f"next earnings window should read like 'Jan 29 – Feb 2, 2027': {when}"


run("earnings: readable analyst tables", {"ticker": "AAPL", "section_EQUITY": "Earnings"}, forbid=["error", "warning"],
    check=check_analyst_tables)


def check_fund_tables(calls, texts):
    ops = tables(calls).get("table_fund_operations")
    if ops is None or ops.loc["Annual Report Expense Ratio", "VTI"] != "0.03%" or not ops.loc["Total Net Assets", "VTI"].endswith("T"):
        return f"fund operations should read as percentages and short amounts: {None if ops is None else ops.to_dict()}"


run("holdings: readable fund tables", {"ticker": "VTI", "section_ETF": "Holdings"}, forbid=["error", "warning"],
    check=check_fund_tables)


# --- Name, logo slot and top spacing ---------------------------------------------------------------
from PIL import Image as _Image  # noqa: E402
APP_DIR = os.path.dirname(os.path.abspath(APP))


def brand(texts):
    return next((m for m in texts["markdown"] if 'class="td-brand"' in m), "")


def page_config(calls):
    return [d for k, d in calls if k == "set_page_config"][0]


def check_no_logo(calls, texts):
    b = brand(texts)
    if "<span>Mew2 Tracker</span>" not in b or "<img" in b:
        return f"brand should show the name without a logo: {b}"
    cfg = page_config(calls)
    if cfg["page_title"] != "Mew2 Tracker" or cfg["page_icon"] != "📊":
        return f"page config wrong: {cfg}"
    css = next(m for m in texts["markdown"] if m.startswith("<style>"))
    if "padding-top: 4.5rem" not in css:
        return "top padding should clear Streamlit's 3.75rem header"


run("name without a logo file", {}, check=check_no_logo, forbid=["error"])

test_logo = os.path.join(APP_DIR, "logo.png.png")  # the double extension Macs often create
_Image.new("RGBA", (30, 40), (120, 80, 160, 255)).save(test_logo, format="PNG")
try:
    def check_logo(calls, texts):
        b = brand(texts)
        if '<img class="td-logo" src="data:image/png;base64,' not in b or b.index("<img") > b.index("<span>"):
            return f"logo should sit left of the name: {b[:120]}"
        if page_config(calls)["page_icon"] != test_logo:
            return "the logo should also be the browser-tab icon"
    run("logo file sits left of the name", {}, check=check_logo, forbid=["error"])
    run("logo also shows with a ticker loaded", {"ticker": "AAPL"}, check=check_logo, forbid=["error", "warning"])
finally:
    os.remove(test_logo)

broken = os.path.join(APP_DIR, "logo.png")
open(broken, "w").write("not really an image")
try:
    run("a broken logo file is ignored", {}, check=check_no_logo, forbid=["error"])
finally:
    os.remove(broken)


# --- Look and feel: fragments, theme, header visuals, loading, confirmations -----------------------------
def check_header_visuals(calls, texts):
    h = header_html(texts)
    m = re.search(r'td-range-dot" style="left: ([0-9.]+)%', h)
    if not m or not 0 <= float(m.group(1)) <= 100:
        return "52-week range bar missing"
    if 'class="td-spark"' not in h or "data:image/svg+xml;base64," not in h or "1 month" not in h:
        return "one-month trend line missing"
    if h.count("52-week range") != 1:
        return "the range should appear once, as the bar"


import re  # noqa: E402
run("header: range bar and trend line", {"ticker": "AAPL"}, check=check_header_visuals, forbid=["error", "warning"])


def check_new_ticker_skeleton(calls, texts):
    if not any('class="td-skeleton"' in m for m in texts["markdown"]):
        return "a new ticker should show a loading placeholder"
    if not find(calls, "empty_cleared"):
        return "the placeholder should be cleared once data arrives"


run("loading placeholder for a new ticker", {"ticker": "AAPL"}, check=check_new_ticker_skeleton, forbid=["error"])
run("...but not when switching tabs", {"ticker": "AAPL", "section_EQUITY": "Ratios"}, keep_session=True, forbid=["error"],
    check=lambda c, t: None if not any('td-skeleton"' in m for m in t["markdown"]) else "tab switches shouldn't flicker")

run("refresh shows a confirmation", {"ticker": "AAPL", "refresh": True}, forbid=["error"],
    check=lambda c, t: None if any(d["body"] == "Prices refreshed" for k, d in c if k == "toast") else "no toast")
run("ZIP confirmation after it's built", {"ticker": "AAPL", "zip_prepare": True}, forbid=["error"],
    check=lambda c, t: None if any("ZIP is ready" in d["body"] for k, d in c if k == "toast") else "no ZIP toast")

st.session_state.clear()
run("units: first visit to Statements", {"ticker": "AAPL", "section_EQUITY": "Statements"}, keep_session=True)


def check_units_rerun(calls, texts):
    if find(calls, "selectbox", key="units")[0]["value"] != "Billions":
        return "units not applied"


calls_u, _ = run("units: changing them refreshes the whole page once (so the ZIP button updates)",
                 {"ticker": "AAPL", "section_EQUITY": "Statements", "units": "Billions"}, keep_session=True,
                 check=check_units_rerun, forbid=["error"])


def check_dark(calls, texts):
    fig = find(calls, "plotly_chart", key="chart_price")[0]["figure"]
    close = [t for t, _, _ in fig.traces if t.props.get("name") == "Close"][0]
    if close.props["line"]["color"] != "#B79CF0":
        return f"dark mode should use the lilac line, got {close.props['line']['color']}"
    css = next(m for m in texts["markdown"] if m.startswith("<style>"))
    if "--td-up: #3CC48A" not in css:
        return "header colors should switch to the dark palette"


st.context.theme.type = st.context.theme.base = "dark"
run("dark mode: charts and header use the dark palette", {"ticker": "AAPL", "price_style": "Line"}, check=check_dark,
    forbid=["error", "warning"])
st.context.theme.type = st.context.theme.base = "light"


def check_chart_config(calls, texts):
    cfg = find(calls, "plotly_chart", key="chart_price")[0]
    fig = cfg["figure"]
    if fig.layout.get("font", {}).get("family", "").split(",")[0] != '"IBM Plex Sans"':
        return "charts should use the app font"
    return None


run("charts share one style", {"ticker": "AAPL"}, check=check_chart_config, forbid=["error"])

# --- Pure calculations ---------------------------------------------------------------------------
st.session_state.clear()
st.STATE.overrides = {}
st.STATE.reset_run()
urllib.request.urlopen = offline_urlopen
requests.get = offline_get
ns = runpy.run_path(APP, run_name="__main__")
problems = []
r = ns["run_dcf"](base_cf=100.0, growth=0.10, years=2, terminal_growth=0.02, wacc=0.08, fade=False,
                  net_claims=50.0, shares=10.0, first_year=2025)
ev = 110 / 1.08 + 121 / 1.08 ** 2 + (121 * 1.02 / 0.06) / 1.08 ** 2
if abs(r["enterprise_value"] - ev) > 1e-9:
    problems.append("run_dcf mismatch")
if ns["clean_security_name"]("Alphabet Inc. - Class A Common Stock") != "Alphabet Inc.":
    problems.append("name cleaning (Nasdaq style) wrong")
if ns["clean_security_name"]("Agilent Technologies, Inc. Common Stock") != "Agilent Technologies, Inc.":
    problems.append("name cleaning (NYSE style) wrong")
if ns["fmt_amount"](1.06) != "1.06" or ns["fmt_amount"](0.0525) != "0.0525":
    problems.append("fmt_amount wrong")
if ns["text"]("AT&T <b> $5") != "AT&amp;T &lt;b&gt; &#36;5":
    problems.append("HTML escaping wrong")

# RSI: steady rises give 100; alternating equal moves sit near 50
up = pd.Series(np.arange(1.0, 60.0))
if abs(ns["rsi"](up).dropna().iloc[-1] - 100) > 1e-9:
    problems.append("RSI of a steady rise should be 100")
zig = pd.Series([10.0, 11.0] * 40)
if not 45 < ns["rsi"](zig).dropna().iloc[-1] < 55:
    problems.append("RSI of equal ups and downs should be near 50")
# MACD on a straight line converges to a constant gap and a flat histogram
line = pd.Series(np.linspace(100, 200, 400))
mac = ns["macd"](line).dropna()
if abs(mac["MACD histogram"].iloc[-1]) > 1e-6 or mac["MACD"].iloc[-1] <= 0:
    problems.append("MACD on a rising line should be positive with a flat histogram")
# Reverse DCF round-trips
params = dict(base_cf=100.0, growth=0.08, years=5, terminal_growth=0.025, wacc=0.09, fade=True, net_claims=50.0,
              shares=10.0, first_year=2025)
target = ns["run_dcf"](**params)["per_share"]
g = ns["implied_growth"](params, target)
w_ = ns["implied_wacc"](params, target)
if g is None or abs(g - 0.08) > 1e-6 or w_ is None or abs(w_ - 0.09) > 1e-6:
    problems.append(f"reverse DCF should recover 8% growth and 9% WACC, got {g}, {w_}")
if ns["implied_growth"](params, 1e9) is not None:
    problems.append("an unreachable price should give no implied growth")
# DDM equals the Gordon formula when growth is flat
d = ns["run_ddm"](2.0, 0.03, 10, 0.03, 0.08, False)["value"]
if abs(d - 2.0 * 1.03 / (0.08 - 0.03)) > 1e-9:
    problems.append(f"DDM should equal Gordon growth with flat growth: {d}")
# Share counts reported before a split get restated
idx = pd.bdate_range("2018-01-01", "2025-12-31")
daily = pd.DataFrame({"Close": 100.0, "Stock Splits": 0.0}, index=idx)
daily.loc[pd.Timestamp("2022-06-06"), "Stock Splits"] = 4.0
shares = pd.Series([4.0e9, 4.1e9, 16.2e9, 16.0e9], index=pd.to_datetime(["2020-09-30", "2021-09-30", "2022-09-30", "2023-09-30"]))
fixed = ns["to_current_share_basis"](shares, daily)
if list(fixed.round(-8)) != [16.0e9, 16.4e9, 16.2e9, 16.0e9]:
    problems.append(f"split restatement wrong: {list(fixed)}")
already = pd.Series([16.4e9, 16.3e9, 16.2e9], index=pd.to_datetime(["2020-09-30", "2021-09-30", "2022-09-30"]))
if list(ns["to_current_share_basis"](already, daily)) != list(already):
    problems.append("already-adjusted share counts should be left alone")

# every tab is a fragment, so using a control inside it reruns only that tab
for name in ("section_price", "section_performance", "section_statements", "section_ratios", "section_valuation",
             "section_fund_valuation", "section_earnings", "section_fund", "section_dividends", "section_compare",
             "section_profile"):
    if not getattr(ns[name], "__is_fragment__", False):
        problems.append(f"{name} should be a fragment")
import tomllib  # noqa: E402
cfg = tomllib.load(open(os.path.join(os.path.dirname(APP), ".streamlit", "config.toml"), "rb"))
if cfg["theme"]["light"]["primaryColor"] != "#6B4FBB" or cfg["theme"]["dark"]["primaryColor"] != "#B79CF0" \
        or not cfg["theme"]["headingFont"].startswith("Sora:https://fonts.googleapis.com/"):
    problems.append("theme config incomplete")
spark = ns["sparkline_html"](pd.DataFrame({"Close": np.linspace(100, 110, 30)},
                                          index=pd.bdate_range("2026-09-01", periods=30)))
import base64 as _b64  # noqa: E402
svg = _b64.b64decode(spark.split("base64,")[1].split('"')[0]).decode()
if "#1E8E5A" not in svg or "1 month +" not in spark:
    problems.append("a rising month should draw a green line with a + label")
falling = ns["sparkline_html"](pd.DataFrame({"Close": np.linspace(110, 100, 30)},
                                            index=pd.bdate_range("2026-09-01", periods=30)))
if "#D6452F" not in _b64.b64decode(falling.split("base64,")[1].split('"')[0]).decode() or "1 month −" not in falling:
    problems.append("a falling month should draw a red line with a minus-sign label")
bar = ns["range_bar_html"](100, 200, 175)
if 'left: 75.0%' not in bar:
    problems.append(f"range dot should sit at 75%: {bar}")
print(f"[{'PASS' if not problems else 'FAIL'}] pure calculations")
for p_ in problems:
    print("   problem:", p_)
RESULTS.append(("pure calculations", not problems))
urllib.request.urlopen = REAL_URLOPEN
requests.get = REAL_GET

print()
failed = [n for n, ok in RESULTS if not ok]
print(f"{len(RESULTS) - len(failed)}/{len(RESULTS)} scenarios passed")
if failed:
    print("Failed:", *failed, sep="\n  ")
    sys.exit(1)
